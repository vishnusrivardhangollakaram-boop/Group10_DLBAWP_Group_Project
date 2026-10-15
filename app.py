"""
LegalCheck India - legal-issue validity check and preliminary complaint drafting
================================================================================
ISOM5240 Deep Learning Business Applications with Python - Group 10 Project
Gollakaram Vishnu Srivardhan (21401910, group leader) and Rahul Singh (2610gnam02)
Proposed for: Karavadi & Associates (https://karavadi.in/) - non-commercial coursework.

User flow:
    1. The client or legal assistant types an English description of the problem (or picks an example).
    2. Pipeline 1  TEXT -> VALIDITY : our fine-tuned Hugging Face text-classification model decides whether the
                                      narrative is a legally valid dispute / issue in India (VALID / NOT_VALID).
    3a. NOT_VALID : the app explains that no legal dispute was recognised, shows the model score, and offers
                    a "New query" button.
    3b. VALID     : Pipeline 2  TEXT -> DRAFT : a pre-trained Hugging Face text-generation model writes the facts,
                    grounds and relief of a preliminary complaint. Python wraps it in a fixed template with
                    __ [Field] placeholders, shows it in an editable box with a missing-information checklist,
                    and exports the CURRENT edited text as a Word (.docx) file.

Model selection (3 candidates per pipeline, see notebooks and Experimental_results.xlsx):
    Pipeline 1 (fine-tuned)  : distilbert/distilbert-base-uncased, google-bert/bert-base-uncased, law-ai/InLegalBERT
    Pipeline 2 (pre-trained) : Qwen/Qwen2.5-0.5B-Instruct, Qwen/Qwen3-0.6B, HuggingFaceTB/SmolLM2-360M-Instruct

Code structure (course conventions):
    IMPORT PART   : libraries, then settings/constants in one place.
    FUNCTION PART : small single-purpose functions (no classes), each with a docstring:
                    1 page style, 2 models, 3 pipeline steps, 4 Word export, 5 screen sections.
    MAIN PART     : main() runs INPUT -> PROCESS -> OUTPUT; started by `if __name__ == "__main__"`.
    Shared logic (decision rule, prompt, template, checks) lives in legalcheck_core.py and is copied
    verbatim into the Colab notebooks, so the evaluated logic and the deployed logic are identical.

Acknowledgement: generative AI (Claude) was used as a coding assistant, as permitted by the course syllabus.
"""

# ==============================
# IMPORT PART
# ==============================
import io
import logging
import threading
import time
from datetime import date

import streamlit as st
import streamlit.components.v1 as components  # tiny HTML/JS snippet used only for smooth auto-scroll
import torch
from docx import Document
from docx.shared import Pt
from transformers import TextIteratorStreamer, pipeline

import legalcheck_core as core

# ---- Settings and constants ----
APP_NAME = "LegalCheck India"
APP_TAGLINE = "Is it a legal dispute in India? If yes, get a preliminary complaint draft in minutes."
COMPANY_NAME = "Karavadi & Associates"
COMPANY_URL = "https://karavadi.in/"
# Fine-tuned in Notebook 1 and uploaded to Hugging Face (change here if the repository name changes).
CLASSIFIER_MODEL_NAME = "VishnuGollakaram/legalcheck-india-validity"
# Selected in Notebook 2 (best pass rate for its runtime on the drafting test cases).
DRAFT_MODEL_NAME = "Qwen/Qwen3-0.6B"
DEFAULT_CPU_THREADS = 2                      # Streamlit Community Cloud gives each app about 2 CPU cores
STREAM_TIMEOUT_SECONDS = 240
LOGGER = logging.getLogger(APP_NAME)
if not LOGGER.handlers:
    logging.basicConfig(level=logging.INFO, format="%(asctime)s | %(levelname)s | %(name)s | %(message)s")

EXAMPLE_NARRATIVES = {
    "Choose an example...": "",
    "Online fraud (valid)": ("Someone called me last Friday pretending to be from my bank in Vijayawada and asked "
                             "for an OTP. After I shared it, Rs 48,000 was debited from my savings account."),
    "Deposit not returned (valid)": ("I vacated my rented flat in Hyderabad on 30 June 2025. The landlord, "
                                     "Mr. Rao, is refusing to return my security deposit of Rs 60,000 although I "
                                     "left the flat in good condition."),
    "Friend's wedding (not valid)": ("My college friend did not invite me to his wedding last month and I feel "
                                     "insulted. I want to take action against him."),
    "Haircut joke (not valid)": ("My friend did not like my new haircut and laughed at it in front of our classmates. "
                                 "I feel embarrassed and want to take action."),
}

PAGE_CSS = """
<style>
/* ---- Karavadi & Associates inspired theme: deep walnut, terracotta red, warm ivory ---- */
@import url('https://fonts.googleapis.com/css2?family=Poppins:wght@500;600;700&family=Source+Sans+3:wght@400;600&display=swap');
:root{--ink:#2b1d1a;--muted:#6d5d57;--line:#e8ddd6;--accent:#b5452f;--deep:#2b1d1a;--gold:#c0503a;
      --ok:#2f6b4f;--okbg:#f1f7f2;--warn:#a5462f;--warnbg:#fbf1ec;--paper:#ffffff;--canvas:#faf7f4}
.stApp{background:var(--canvas);color:var(--ink);font-family:'Source Sans 3',sans-serif}
[data-testid="stHeader"]{background:transparent}
[data-testid="stSidebar"]{background:#2b1d1a}
[data-testid="stSidebar"] *{color:#f3e9e3!important}
[data-testid="stSidebar"] a{color:#f0a48f!important}
[data-testid="stSidebar"] code{background:transparent!important;color:#f0a48f!important;padding:0!important;
                              font-size:13.5px;border:none!important}   /* model names: no white box on dark sidebar */
[data-testid="stMainBlockContainer"]{max-width:1180px;padding-top:2.0rem}
h1,h2,h3,h4{font-family:'Poppins',sans-serif;letter-spacing:-.01em;color:var(--ink)}
.brand{display:flex;align-items:center;gap:16px;padding:22px 26px;margin-bottom:18px;border-radius:18px;
       background:linear-gradient(120deg,#2b1d1a 0%,#4a2a22 60%,#7a3424 100%);box-shadow:0 10px 30px #2b1d1a33}
.brand-mark{width:54px;height:54px;border-radius:14px;background:#c0503a;color:#fff;display:flex;
            align-items:center;justify-content:center;font-size:28px;box-shadow:0 4px 14px #0004}
.brand-name{font-family:'Poppins',sans-serif;font-size:32px;font-weight:700;color:#fff;line-height:1.05}
.brand-name span{color:#f0a48f}
.brand-sub{font-size:14.5px;color:#eadbd3;margin-top:6px}
.steps{display:flex;gap:10px;margin:4px 0 18px 0;flex-wrap:wrap}
.step{padding:7px 14px;border-radius:999px;border:1px solid var(--line);background:#fff;font-size:13.5px;color:var(--muted)}
.step.on{background:var(--accent);border-color:var(--accent);color:#fff;box-shadow:0 3px 10px #b5452f40}
.card{background:var(--paper);border:1px solid var(--line);border-radius:16px;padding:18px 20px;margin:8px 0 14px 0}
.verdict-ok{background:var(--okbg);border:1px solid #cfe3d6;border-left:7px solid var(--ok);border-radius:16px;
            padding:18px 20px;box-shadow:0 6px 18px #2f6b4f1a}
.verdict-no{background:var(--warnbg);border:1px solid #f0d3c7;border-left:7px solid var(--warn);border-radius:16px;
            padding:18px 20px;box-shadow:0 6px 18px #a5462f1a}
.verdict-title{font-family:'Poppins',sans-serif;font-size:21px;font-weight:700;margin-bottom:6px}
.meter{height:9px;background:#efe6e0;border-radius:6px;overflow:hidden;margin-top:12px}
.meter>div{height:9px;border-radius:6px}
.small{font-size:13px;color:var(--muted);margin-top:6px}
.chip{display:inline-block;padding:4px 11px;margin:3px 5px 3px 0;border-radius:999px;background:#fbefe9;
      border:1px solid #efcfc2;font-size:13px;color:#7a3424}
textarea{border-radius:12px!important;border-color:var(--line)!important;font-size:16px!important}
button[kind="primary"]{background:var(--accent)!important;border-color:var(--accent)!important;border-radius:10px!important;
                       font-weight:600!important}
button[kind="primary"]:hover{background:#2b1d1a!important;border-color:#2b1d1a!important}
button[kind="secondary"]{border-radius:10px!important;border-color:#d9c6bc!important}
[data-testid="stDownloadButton"] button{border-radius:10px!important}
</style>
"""


# ==============================
# FUNCTION PART
# ==============================
# ---- 1 Page style ----
def apply_page_style():
    """Set the page configuration and inject the CSS theme."""
    st.set_page_config(page_title=f"{APP_NAME} | Legal issue check & draft", page_icon="⚖️", layout="wide")
    st.markdown(PAGE_CSS, unsafe_allow_html=True)


def show_header():
    """Show the brand header and one-sentence purpose."""
    st.markdown(
        f"""<div class="brand"><div class="brand-mark">⚖️</div><div>
        <div class="brand-name">LegalCheck <span>India</span></div>
        <div class="brand-sub">{APP_TAGLINE}<br>Prototype proposed for {COMPANY_NAME} · coursework, not legal advice</div>
        </div></div>""",
        unsafe_allow_html=True,
    )


def show_steps(active_step):
    """Show the three-step progress bar (1 Describe, 2 Check, 3 Draft & download)."""
    labels = ["1 · Describe the issue", "2 · Validity check", "3 · Review & download draft"]
    chips = "".join(f'<span class="step{" on" if index + 1 <= active_step else ""}">{label}</span>'
                    for index, label in enumerate(labels))
    st.markdown(f'<div class="steps">{chips}</div>', unsafe_allow_html=True)


def auto_scroll(anchor_id):
    """Smoothly scroll the page to an HTML anchor so the newest result is on screen without manual scrolling."""
    LOGGER.info(f"UI auto-scroll to #{anchor_id}")
    components.html(
        f"""<script>
        const target = window.parent.document.getElementById("{anchor_id}");
        if (target) {{ setTimeout(() => target.scrollIntoView({{behavior: "smooth", block: "start"}}), 250); }}
        </script>""",
        height=0,
    )


# ---- 2 Models ----
@st.cache_resource(show_spinner=False)
def load_classifier(model_name):
    """Load Pipeline 1 (fine-tuned text-classification) once per server."""
    torch.set_num_threads(DEFAULT_CPU_THREADS)
    start_time = time.perf_counter()
    classifier = pipeline(task="text-classification", model=model_name, device=-1)
    LOGGER.info(f"Classifier loaded in {time.perf_counter() - start_time:.2f} s")
    return classifier


@st.cache_resource(show_spinner=False)
def load_generator(model_name):
    """Load Pipeline 2 (pre-trained text-generation) once per server, on first VALID case."""
    torch.set_num_threads(DEFAULT_CPU_THREADS)
    start_time = time.perf_counter()
    generator = pipeline(task="text-generation", model=model_name, device=-1, torch_dtype=torch.bfloat16)  # half the memory of float32
    LOGGER.info(f"Generator loaded in {time.perf_counter() - start_time:.2f} s")
    return generator


# ---- 3 Pipeline steps ----
def reset_results():
    """Forget the previous decision and draft (used by New query and when the narrative changes)."""
    for key in ["decision", "valid_score", "classify_seconds", "draft_text", "draft_seconds", "checked_text"]:
        st.session_state.pop(key, None)


def run_validity_check(narrative):
    """Validate the input, then run Pipeline 1 and store the decision in session state."""
    is_ok, message = core.validate_narrative(narrative=narrative)
    if not is_ok:
        st.warning(message)
        return
    try:
        with st.spinner("Checking whether this is a legal dispute in India..."):
            classifier = load_classifier(model_name=CLASSIFIER_MODEL_NAME)
            decision, valid_score, seconds = core.classify_narrative(classifier=classifier, narrative=narrative)
    except Exception as error:                      # visible failure, never a silent fallback
        LOGGER.exception("Classifier failed")
        st.error(f"The validity model could not run ({type(error).__name__}). Please try again in a minute.")
        return
    st.session_state.update(decision=decision, valid_score=valid_score, classify_seconds=seconds,
                            checked_text=narrative)
    st.session_state.pop("draft_text", None)
    st.session_state["scroll_to"] = "verdict-anchor"         # show the verdict without manual scrolling
    LOGGER.info(f"TIMING classify={seconds:.2f}s decision={decision} score={valid_score:.3f}")


def stream_draft(narrative):
    """Run Pipeline 2 with live token streaming; return (complaint_text, seconds) or (None, 0) on failure."""
    try:
        with st.spinner("Loading the drafting model (first use only)..."):
            generator = load_generator(model_name=DRAFT_MODEL_NAME)
        streamer = TextIteratorStreamer(generator.tokenizer, skip_prompt=True, skip_special_tokens=True,
                                        timeout=STREAM_TIMEOUT_SECONDS)
        result = {}

        def worker():
            """Generate in the background so the words can appear on screen while they are written."""
            try:
                result["body"], result["seconds"] = core.generate_draft_body(generator=generator, narrative=narrative,
                                                                             streamer=streamer)
            except Exception as error:              # passed back to the main thread and shown to the user
                result["error"] = error
                streamer.end()

        thread = threading.Thread(target=worker, daemon=True)
        thread.start()
        with st.container(border=True):
            st.caption("Writing the facts, grounds and relief...")
            st.write_stream(streamer)
        thread.join(timeout=STREAM_TIMEOUT_SECONDS)
        if "error" in result or "body" not in result:
            raise RuntimeError(str(result.get("error", "generation did not finish")))
    except Exception as error:
        LOGGER.exception("Drafting failed")
        st.error(f"The drafting model could not finish ({type(error).__name__}). Please press Generate again.")
        return None, 0.0
    checks = core.score_draft_body(body_text=result["body"], key_facts=[])
    if not checks["structure_ok"]:
        st.warning("The draft structure looks incomplete - please review the text carefully before use.")
    LOGGER.info(f"TIMING draft={result['seconds']:.2f}s words={checks['words']}")
    return core.assemble_complaint(body_text=result["body"]), result["seconds"]


# ---- 4 Word export ----
def build_docx_bytes(draft_text):
    """Create an editable Word document from the CURRENT (possibly edited) draft text."""
    LOGGER.info(f"Building Word export ({core.count_words(text=draft_text)} words)")
    document = Document()
    style = document.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)
    for line in draft_text.split("\n"):
        stripped = line.strip()
        if not stripped:
            continue
        paragraph = document.add_paragraph()
        is_heading = stripped.isupper() or (stripped.split(":")[0].isupper() and ":" in stripped[:40])
        run = paragraph.add_run(stripped)
        run.bold = is_heading and len(stripped) < 120
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


# ---- 5 Screen sections ----
def show_sidebar():
    """Sidebar: how it works, models used and disclaimer."""
    with st.sidebar:
        st.markdown("### How it works")
        st.markdown("1. Describe the problem in plain English.\n"
                    "2. **Fine-tuned model** checks if it is a legally valid dispute / issue in India.\n"
                    "3. If valid, a **drafting model** prepares a preliminary complaint you can edit and download.")
        st.markdown("### Models")
        st.markdown(f"- Validity: [`{CLASSIFIER_MODEL_NAME}`](https://huggingface.co/{CLASSIFIER_MODEL_NAME})\n"
                    f"- Drafting: [`{DRAFT_MODEL_NAME}`](https://huggingface.co/{DRAFT_MODEL_NAME})")
        st.markdown("### Important")
        st.caption("Coursework prototype (ISOM5240, Group 10) proposed for "
                   f"[{COMPANY_NAME}]({COMPANY_URL}). Output is a preliminary draft for advocate review, "
                   "not legal advice. Do not enter real personal data.")


def get_narrative_input():
    """INPUT: example picker + narrative text box. Returns (narrative, check_clicked)."""
    st.markdown("#### Describe the issue")
    example_name = st.selectbox("Try an example", options=list(EXAMPLE_NARRATIVES.keys()), key="example_name")
    if example_name != st.session_state.get("last_example") and EXAMPLE_NARRATIVES.get(example_name):
        st.session_state["narrative"] = EXAMPLE_NARRATIVES[example_name]
        st.session_state["last_example"] = example_name
    narrative = st.text_area("What happened? Include who, what, when, where and any amounts.", key="narrative",
                             height=170, placeholder="e.g. My landlord in Pune is not returning my Rs 40,000 deposit...")
    word_count = core.count_words(text=narrative)
    st.caption(f"{word_count} words · minimum {core.MINIMUM_INPUT_WORDS}, maximum {core.MAXIMUM_INPUT_WORDS}")
    if st.session_state.get("checked_text") and narrative != st.session_state.get("checked_text"):
        reset_results()                              # edited narrative -> previous result is stale
    left, right = st.columns([1, 1])
    check_clicked = left.button("Check validity", type="primary", use_container_width=True)
    right.button("New query", on_click=start_new_query, use_container_width=True, key="new_query_top")
    return narrative, check_clicked


def start_new_query():
    """Clear the narrative and every result so the user can start again."""
    LOGGER.info("New query: clearing narrative, verdict and draft")
    reset_results()
    st.session_state["narrative"] = ""
    st.session_state["example_name"] = "Choose an example..."
    st.session_state["last_example"] = "Choose an example..."


def show_verdict():
    """OUTPUT 1: the VALID / NOT_VALID verdict card with the model score."""
    decision = st.session_state["decision"]
    valid_score = st.session_state["valid_score"]
    seconds = st.session_state["classify_seconds"]
    if decision == core.VALID_LABEL:
        css_class, title, colour = "verdict-ok", "✅ Legally valid dispute / issue in India", "#2f6b4f"
        text = "The fine-tuned model recognises an actionable legal grievance. A preliminary complaint can be drafted."
    else:
        css_class, title, colour = "verdict-no", "⚠️ Not a legally valid dispute / issue in India", "#a5462f"
        text = ("The model did not recognise a legal wrong that Indian law would act on (for example a personal, "
                "social or matter-of-taste grievance, a lawful act, an unenforceable claim such as a bet, or an "
                "event with no Indian connection). No complaint is drafted. You may rephrase with more facts or "
                "start a new query.")
    st.markdown('<div id="verdict-anchor"></div>', unsafe_allow_html=True)   # target for auto-scroll
    st.markdown(
        f"""<div class="{css_class}"><div class="verdict-title">{title}</div><div>{text}</div>
        <div class="meter"><div style="width:{valid_score * 100:.0f}%;background:{colour}"></div></div>
        <div class="small">Model score for VALID: {valid_score:.2f} (threshold {core.VALIDITY_THRESHOLD:.2f}) ·
        checked in {seconds:.2f} s</div></div>""",
        unsafe_allow_html=True,
    )


def show_draft_section(narrative):
    """OUTPUT 2 (VALID only): generate, edit, check missing fields and download the complaint."""
    st.markdown('<div id="draft-anchor"></div>', unsafe_allow_html=True)     # target for auto-scroll
    st.markdown("#### Preliminary complaint draft")
    if "draft_text" not in st.session_state:
        if st.button("Generate complaint draft", type="primary"):
            complaint_text, seconds = stream_draft(narrative=narrative)
            if complaint_text:
                st.session_state.update(draft_text=complaint_text, draft_seconds=seconds,
                                        draft_editor=complaint_text, scroll_to="draft-anchor")
                LOGGER.info(f"Draft stored in session ({core.count_words(text=complaint_text)} words)")
                st.rerun()
        st.caption("Drafting runs on a CPU and usually takes 30-90 seconds.")
        return
    st.caption(f"Drafted in {st.session_state['draft_seconds']:.2f} s · edit freely; the download uses your edits.")
    edited_text = st.text_area("Editable draft", key="draft_editor", height=480, label_visibility="collapsed")
    missing_fields = core.list_missing_fields(draft_text=edited_text)
    st.markdown("**Missing information to complete** " + (
        "".join(f'<span class="chip">{field}</span>' for field in missing_fields) if missing_fields
        else '<span class="chip">None - all placeholders filled</span>'), unsafe_allow_html=True)
    left, middle, right = st.columns(3)
    left.download_button("Download Word (.docx)", data=build_docx_bytes(draft_text=edited_text),
                         file_name=f"preliminary_complaint_{date.today():%Y%m%d}.docx",
                         mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                         use_container_width=True, type="primary")
    middle.download_button("Download text (.txt)", data=edited_text.encode("utf-8"),
                           file_name=f"preliminary_complaint_{date.today():%Y%m%d}.txt", use_container_width=True)
    right.button("New query", on_click=start_new_query, use_container_width=True, key="new_query_bottom")


# ==============================
# MAIN PART
# ==============================
def main():
    """INPUT -> PROCESS -> OUTPUT (course structure). Streamlit re-runs this function after every user action."""
    # ---- Step 1: page set-up (theme, header, sidebar, progress chips) ----
    apply_page_style()
    show_header()
    show_sidebar()
    active_step = 3 if st.session_state.get("draft_text") else 2 if st.session_state.get("decision") else 1
    show_steps(active_step=active_step)
    LOGGER.info(f"Page render | step={active_step} | decision={st.session_state.get('decision')}")

    # ---- Step 2: INPUT - narrative text box, example picker and buttons ----
    narrative, check_clicked = get_narrative_input()

    # ---- Step 3: PROCESS - Pipeline 1 (fine-tuned validity classifier) when the user presses Check validity ----
    if check_clicked:
        LOGGER.info(f"Check validity pressed ({core.count_words(text=narrative)} words)")
        run_validity_check(narrative=narrative)

    # ---- Step 4: OUTPUT - verdict card; VALID -> Pipeline 2 drafting section, NOT_VALID -> new-query option ----
    if st.session_state.get("decision"):
        show_verdict()
        if st.session_state["decision"] == core.VALID_LABEL:
            show_draft_section(narrative=st.session_state["checked_text"])
        else:
            st.button("Start a new query", on_click=start_new_query, type="primary", key="new_query_invalid")

    # ---- Step 5: one-time smooth auto-scroll to the newest output (verdict or finished draft) ----
    scroll_target = st.session_state.pop("scroll_to", None)
    if scroll_target:
        auto_scroll(anchor_id=scroll_target)


if __name__ == "__main__":
    main()
